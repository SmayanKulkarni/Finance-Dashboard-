"""
Build the Bank Loan Report Excel dashboard (Summary / Overview / Details).

Excel re-creation of the Power BI "Bank Loan Analysis" project
(https://youtu.be/3I8wd1AShXs) using data/financial_loan.csv.

Everything in the workbook is driven by live Excel formulas, so the
dropdown "slicers" recalculate every KPI, chart and grid inside Excel.

    pip install pandas xlsxwriter pillow
    python build_dashboard.py            ->  Bank_Loan_Dashboard.xlsx
"""
import datetime as dt
import os
import shutil
import subprocess
import sys
import tempfile

import pandas as pd
import xlsxwriter
from xlsxwriter.utility import xl_rowcol_to_cell as rc, xl_col_to_name

HERE = os.path.dirname(os.path.abspath(__file__))
CSV = os.path.join(HERE, "data", "financial_loan.csv")
OUT = os.environ.get("BLD_OUT", os.path.join(HERE, "Bank_Loan_Dashboard.xlsx"))
LOGO = os.path.join(HERE, "assets", "bank_logo.png")
CACHE_FILE = os.environ.get("BLD_CACHE")  # recalculated copy used to pre-fill formula results


def main_two_pass():
    """Build once, let LibreOffice calculate it, then rebuild with every formula's
    cached result filled in. Without cached values Excel shows zeros in Protected
    View (before 'Enable Editing'), and chart labels/titles start out blank."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found - building without cached values "
              "(Excel recalculates on open).")
        return False
    tmp = tempfile.mkdtemp()
    prof = os.path.join(tmp, "profile", "user")
    os.makedirs(prof)
    with open(os.path.join(prof, "registrymodifications.xcu"), "w") as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?><oor:items '
                'xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
                '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>'
                '</item></oor:items>')
    first = os.path.join(tmp, "pass1.xlsx")
    env = dict(os.environ, BLD_OUT=first, BLD_NO_RECALC="1")
    subprocess.run([sys.executable, __file__], env=env, check=True)
    subprocess.run([soffice, f"-env:UserInstallation=file://{os.path.dirname(prof)}",
                    "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                    "--outdir", os.path.join(tmp, "calc"), first],
                   check=True, capture_output=True, timeout=1200)
    env = dict(os.environ, BLD_NO_RECALC="1",
               BLD_CACHE=os.path.join(tmp, "calc", "pass1.xlsx"))
    subprocess.run([sys.executable, __file__], env=env, check=True)
    shutil.rmtree(tmp, ignore_errors=True)
    return True


if __name__ == "__main__" and not os.environ.get("BLD_NO_RECALC"):
    if main_two_pass():
        sys.exit(0)

# ----------------------------------------------------------------------------
# Theme (sampled from the Power BI report)
# ----------------------------------------------------------------------------
C = dict(
    header="#050A14", panel="#08111F", page="#0C1A3D", card="#0B2D5E",
    card_hi="#9C4F2E", text="#FFFFFF", sub="#B9C7E6", blue="#3DA5FF",
    gold="#E8C13A", pink="#F2A7DA", mtd="#F4B860", magenta="#E000E6",
    grey="#5F5F5F", btn="#333333", btn_on="#E6E6E6", dd="#0B2D5E",
    line="#1E90FF", area="#0D55B5", lav="#A98BE8", purp="#C83CF0",
    t36="#6FB6FF", t60="#6236FF", row1="#3B2A73", row2="#262626",
    map_lo="#F7D9C8", map_hi="#A3421B", map_na="#D0D0D0",
)
HOME_COLORS = ["#1E90FF", "#16167A", "#E8743B", "#C83CF0", "#E8C13A"]
FONT = "Segoe UI"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MEASURES = ["Total Loan Applications", "Total Funded Amount", "Total Amount Received"]
EMP_ORDER = ["< 1 year", "1 year", "2 years", "3 years", "4 years", "5 years",
             "6 years", "7 years", "8 years", "9 years", "10+ years"]

# US tile-grid map (col, row) -- stands in for the Power BI filled map.
TILES = {
    "AK": (0, 0), "ME": (10, 0), "VT": (9, 1), "NH": (10, 1),
    "WA": (0, 2), "ID": (1, 2), "MT": (2, 2), "ND": (3, 2), "MN": (4, 2), "IL": (5, 2),
    "WI": (6, 2), "MI": (7, 2), "NY": (8, 2), "RI": (9, 2), "MA": (10, 2),
    "OR": (0, 3), "NV": (1, 3), "WY": (2, 3), "SD": (3, 3), "IA": (4, 3), "IN": (5, 3),
    "OH": (6, 3), "PA": (7, 3), "NJ": (8, 3), "CT": (9, 3),
    "CA": (0, 4), "UT": (1, 4), "CO": (2, 4), "NE": (3, 4), "MO": (4, 4), "KY": (5, 4),
    "WV": (6, 4), "VA": (7, 4), "MD": (8, 4), "DE": (9, 4),
    "AZ": (1, 5), "NM": (2, 5), "KS": (3, 5), "AR": (4, 5), "TN": (5, 5), "NC": (6, 5),
    "SC": (7, 5), "DC": (8, 5),
    "OK": (3, 6), "LA": (4, 6), "MS": (5, 6), "AL": (6, 6), "GA": (7, 6),
    "HI": (0, 7), "TX": (3, 7), "FL": (8, 7),
}

# ----------------------------------------------------------------------------
# Data
# ----------------------------------------------------------------------------
df = pd.read_csv(CSV)
for col in ("issue_date", "last_credit_pull_date", "last_payment_date", "next_payment_date"):
    df[col] = pd.to_datetime(df[col], format="%d-%m-%Y")
df["term"] = df["term"].str.strip()
df["good_vs_bad"] = df["loan_status"].map(
    {"Fully Paid": "Good Loan", "Current": "Good Loan", "Charged Off": "Bad Loan"})
df["issue_month"] = df["issue_date"].dt.month
df["issue_month_name"] = df["issue_date"].dt.strftime("%b")
df["year_month"] = df["issue_date"].dt.year * 100 + df["issue_date"].dt.month
df = df.sort_values("id").reset_index(drop=True)

# Columns A:J are exactly the Details grid (FILTER spills A:J).
DATA_COLS = [
    "id", "purpose", "home_ownership", "grade", "sub_grade", "issue_date",
    "loan_amount", "int_rate", "installment", "total_payment",
    "address_state", "emp_length", "emp_title", "term", "loan_status",
    "verification_status", "application_type", "last_credit_pull_date",
    "last_payment_date", "next_payment_date", "member_id", "annual_income",
    "dti", "total_acc", "good_vs_bad", "issue_month", "issue_month_name", "year_month",
]
df = df[DATA_COLS]
N = len(df)
LAST = N + 1  # last Excel row of data

NAMES = {  # defined name -> data column
    "d_id": "id", "d_purpose": "purpose", "d_home": "home_ownership",
    "d_grade": "grade", "d_issue": "issue_date", "d_amount": "loan_amount",
    "d_rate": "int_rate", "d_payment": "total_payment", "d_state": "address_state",
    "d_emp": "emp_length", "d_term": "term", "d_status": "loan_status",
    "d_dti": "dti", "d_gb": "good_vs_bad", "d_mon": "issue_month", "d_ym": "year_month",
}

lists = {
    "Grade": sorted(df.grade.unique()),
    "State": sorted(df.address_state.unique()),
    "Purpose": sorted(df.purpose.unique(), key=str.lower),
    "Term": sorted(df.term.unique()),
    "Employee Length": [e for e in EMP_ORDER if e in set(df.emp_length)],
    "Home Ownership": list(df.home_ownership.value_counts().index),
    "Good v Bad Loan": ["Good Loan", "Bad Loan"],
}

# ----------------------------------------------------------------------------
# Workbook + format cache
# ----------------------------------------------------------------------------
CACHE = {}
if CACHE_FILE:
    import openpyxl
    _cwb = openpyxl.load_workbook(CACHE_FILE, data_only=True)
    for _name in ("Summary", "Overview", "Details", "Calc"):
        for _row in _cwb[_name].iter_rows():
            for _c in _row:
                v = _c.value
                if v is None or (isinstance(v, str) and v.startswith("#")):
                    continue
                if isinstance(v, dt.datetime):
                    v = (v - dt.datetime(1899, 12, 30)).days
                CACHE[(_name, _c.row - 1, _c.column - 1)] = v

_orig_write_formula = xlsxwriter.worksheet.Worksheet._write_formula


def _write_formula_cached(self, row, col, formula, cell_format=None, value=0):
    v = CACHE.get((self.name, row, col))
    if v is not None and value == 0:
        value = ("TRUE" if v else "FALSE") if isinstance(v, bool) else v
    return _orig_write_formula(self, row, col, formula, cell_format, value)


xlsxwriter.worksheet.Worksheet._write_formula = _write_formula_cached

wb = xlsxwriter.Workbook(OUT, {"use_future_functions": True})
wb.set_calc_mode("auto")
wb.set_size(1600, 1000)
_fmt = {}


def F(**kw):
    kw.setdefault("font_name", FONT)
    kw.setdefault("font_size", 10)
    kw.setdefault("valign", "vcenter")
    key = tuple(sorted(kw.items()))
    if key not in _fmt:
        _fmt[key] = wb.add_format(kw)
    return _fmt[key]


NF = dict(
    k='[>=1000]#,##0.0,"K";#,##0',
    m='[>=1000000]$#,##0.0,,"M";[>=1000]$#,##0.0,"K";$#,##0',
    pct2="0.00%", pct1="0.0%", int="#,##0", usd="$#,##0", usd1="$#,##0.0",
)

ws_sum = wb.add_worksheet("Summary")
ws_ovr = wb.add_worksheet("Overview")
ws_det = wb.add_worksheet("Details")
ws_data = wb.add_worksheet("Data")
ws_calc = wb.add_worksheet("Calc")
ws_list = wb.add_worksheet("Lists")
ws_guide = wb.add_worksheet("Guide")

# ----------------------------------------------------------------------------
# Data sheet: Excel Table + defined names
# ----------------------------------------------------------------------------
date_f = F(num_format="dd-mmm-yyyy", valign="top")
col_formats = {
    "issue_date": date_f, "last_credit_pull_date": date_f, "last_payment_date": date_f,
    "next_payment_date": date_f, "int_rate": F(num_format="0.00%", valign="top"),
    "dti": F(num_format="0.00%", valign="top"), "loan_amount": F(num_format=NF["usd"], valign="top"),
    "total_payment": F(num_format=NF["usd"], valign="top"),
    "installment": F(num_format="$#,##0.00", valign="top"),
    "annual_income": F(num_format=NF["usd"], valign="top"),
}
rows = df.astype(object).where(df.notna(), None).values.tolist()
ws_data.add_table(0, 0, N, len(DATA_COLS) - 1, {
    "name": "tblLoan", "style": "Table Style Medium 2", "data": rows,
    "columns": [{"header": c, "format": col_formats.get(c)} for c in DATA_COLS],
})
for i, c in enumerate(DATA_COLS):
    ws_data.set_column(i, i, max(11, min(24, len(c) + 3)))
ws_data.freeze_panes(1, 1)
for nm, col in NAMES.items():
    letter = xl_col_to_name(DATA_COLS.index(col))
    wb.define_name(nm, f"=Data!${letter}$2:${letter}${LAST}")

# ----------------------------------------------------------------------------
# Lists sheet (dropdown sources)
# ----------------------------------------------------------------------------
LIST_REF = {}
ws_list.write(0, 0, "Select Measure", F(bold=True))
for i, m in enumerate(MEASURES):
    ws_list.write(i + 1, 0, m)
LIST_REF["Select Measure"] = f"=Lists!$A$2:$A${len(MEASURES) + 1}"
for j, (name, vals) in enumerate(lists.items(), start=1):
    ws_list.write(0, j, name, F(bold=True))
    ws_list.write(1, j, "All")
    for i, v in enumerate(vals):
        ws_list.write(i + 2, j, v)
    L = xl_col_to_name(j)
    LIST_REF[name] = f"=Lists!${L}$2:${L}${len(vals) + 2}"
ws_list.set_column(0, 8, 22)

# ----------------------------------------------------------------------------
# Page definitions: which slicers each page owns (mirrors the Power BI pages)
# ----------------------------------------------------------------------------
DIM = {  # slicer label -> defined name of data column
    "Good v Bad Loan": "d_gb", "Grade": "d_grade", "State": "d_state",
    "Purpose": "d_purpose", "Term": "d_term", "Employee Length": "d_emp",
    "Home Ownership": "d_home",
}
PAGE_FILTERS = {
    "Summary": ["Purpose", "Grade", "State"],
    "Overview": ["Good v Bad Loan", "Grade", "State",
                 "Term", "Employee Length", "Purpose", "Home Ownership"],
    "Details": ["Good v Bad Loan", "Grade", "State"],
}
SEL = {p: {} for p in PAGE_FILTERS}   # page -> slicer -> absolute cell on dashboard
CRIT = {p: {} for p in PAGE_FILTERS}  # page -> slicer -> criteria cell on Calc


def absref(sheet, r, c):
    return f"{sheet}!{rc(r, c, True, True)}"


# ----------------------------------------------------------------------------
# Calc sheet helper: sequential blocks of formulas
# ----------------------------------------------------------------------------
class Calc:
    def __init__(self, ws):
        self.ws, self.r = ws, 0
        ws.set_column(0, 0, 34)
        ws.set_column(1, 12, 16)
        self.h1 = F(bold=True, font_size=12, font_color="#FFFFFF", bg_color="#0B2D5E")
        self.h2 = F(bold=True, bg_color="#DCE6F5")

    def title(self, text):
        self.r += 1
        self.ws.merge_range(self.r, 0, self.r, 10, text, self.h1)
        self.r += 1

    def header(self, cols):
        for j, c in enumerate(cols):
            self.ws.write(self.r, j, c, self.h2)
        self.r += 1

    def put(self, label, formula, num=None):
        """Write label in A and formula in B; return absolute ref of B."""
        self.ws.write(self.r, 0, label)
        self.ws.write_formula(self.r, 1, formula, F(num_format=num) if num else None)
        ref = absref("Calc", self.r, 1)
        self.r += 1
        return ref


calc = Calc(ws_calc)
calc.ws.write(0, 0, "Calculation engine - every dashboard number comes from here.",
              F(bold=True, font_size=12))


def args(page, exclude=None, extra=()):
    parts = [f"{DIM[s]},{CRIT[page][s]}" for s in PAGE_FILTERS[page] if s != exclude]
    parts += [f"{a},{b}" for a, b in extra]
    return ",".join(parts)


def money_label(v, money):
    """Text label like $54.0M / 4.3K, locale-safe (FIXED)."""
    return (f'IF({money},"$","")&IF(ABS({v})>=1000000,FIXED({v}/1000000,1)&"M",'
            f'IF(ABS({v})>=1000,FIXED({v}/1000,1)&"K",FIXED({v},0)))')


# ----------------------------------------------------------------------------
# Dashboard canvas helpers (Summary / Overview use a 20px grid)
# ----------------------------------------------------------------------------
GRID_COLS, GRID_ROWS = 77, 44


def fill(ws, r1, c1, r2, c2, fmt):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.write_blank(r, c, None, fmt)


def text(ws, r1, c1, r2, c2, value, fmt):
    if (r1, c1) == (r2, c2):
        ws.write(r1, c1, value, fmt)
    else:
        ws.merge_range(r1, c1, r2, c2, "", fmt)
        if isinstance(value, str) and value.startswith("="):
            ws.write_formula(r1, c1, value, fmt)
        else:
            ws.write(r1, c1, value, fmt)


def setup_canvas(ws, page):
    ws.hide_gridlines(2)
    ws.set_zoom(90)
    ws.set_column_pixels(0, 10, 20, F(bg_color=C["panel"]))
    ws.set_column_pixels(11, GRID_COLS - 1, 20, F(bg_color=C["page"]))
    for r in range(GRID_ROWS + 6):
        ws.set_row_pixels(r, 20)
    header(ws, 0, 0, 2, GRID_COLS - 1, page)
    ws.insert_image(3, 2, LOGO, {"x_offset": 12, "y_offset": 4, "x_scale": 0.5, "y_scale": 0.5,
                                 "object_position": 3})


def header(ws, r1, c1, r2, c2, page):
    fmt = F(bg_color=C["header"], font_color=C["text"], bold=True, font_size=22,
            align="center")
    ws.merge_range(r1, c1, r2, c2, "", fmt)
    ws.write_rich_string(r1, c1,
                         F(font_color=C["text"], bold=True, font_size=22), "BANK LOAN REPORT | ",
                         F(font_color=C["blue"], bold=True, font_size=22), page.upper(), fmt)


def nav_buttons(ws, page, rows, c1, c2):
    for name, r in zip(["Summary", "Overview", "Details"], rows):
        on = name == page
        fmt = F(bg_color=C["btn_on"] if on else C["btn"], font_color="#000000" if on else C["text"],
                align="center", font_size=11, bold=on, border=1,
                border_color="#9A9A9A" if on else "#4A4A4A", underline=0)
        r1, r2 = r
        if (r1, c1) != (r2, c2):
            ws.merge_range(r1, c1, r2, c2, "", fmt)
        ws.write_url(r1, c1, f"internal:'{name}'!A1", fmt, name, tip=f"Go to {name}")


def dropdown(ws, page, slicer, r, c1, c2, label_fmt, dd_fmt, label=None):
    text(ws, r, c1, r, c2, label or slicer, label_fmt)
    src = "Select Measure" if slicer == "Select Measure" else slicer
    default = MEASURES[0] if slicer == "Select Measure" else "All"
    text(ws, r + 1, c1, r + 1, c2, default, dd_fmt)
    ws.data_validation(r + 1, c1, r + 1, c2, {
        "validate": "list", "source": LIST_REF[src],
        "input_title": slicer, "input_message": "Pick a value (All = no filter)",
        "error_title": "Invalid selection", "error_message": "Choose a value from the list.",
    })
    ws.write(r + 1, c2 + 1, "▼", F(font_color=C["sub"], font_size=8, align="left",
                                    bg_color=C["panel"]))
    cell = absref(ws.get_name(), r + 1, c1)
    if slicer != "Select Measure":
        SEL[page][slicer] = cell
    return cell


PANEL_LABEL = F(font_color=C["text"], bold=True, font_size=11, bg_color=C["panel"])
PANEL_DD = F(font_color=C["text"], font_size=9, bg_color=C["dd"], border=1,
             border_color="#9FB6E0", indent=1)
PANEL_HINT = F(font_color=C["sub"], italic=True, font_size=8, bg_color=C["panel"], text_wrap=True)


def kpi_card(ws, r0, c0, c1, title, vals, hi=False, gap_borders=False):
    """vals = (value_ref, value_fmt, mtd_ref, mtd_fmt, mom_ref)."""
    bg = C["card_hi"] if hi else C["card"]
    extra = {}
    if gap_borders:
        extra = dict(left=5, right=5, left_color=C["page"], right_color=C["page"])
    for cc in range(c0, c1 + 1):
        edge = {}
        if gap_borders and cc == c0:
            edge.update(left=5, left_color=C["page"])
        if gap_borders and cc == c1:
            edge.update(right=5, right_color=C["page"])
        fill(ws, r0, cc, r0 + 6, cc, F(bg_color=bg, **edge))
    mid = c0 + (c1 - c0 + 1) // 2
    value, vfmt, mtd, mfmt, mom = vals
    text(ws, r0, c0, r0, c1, title,
         F(bg_color=bg, font_color=C["sub"] if not hi else "#F3D3C0", bold=True,
           font_size=11, align="center", **extra))
    text(ws, r0 + 1, c0, r0 + 2, c1, f"={value}",
         F(bg_color=bg, font_color=C["text"], bold=True, font_size=24, align="center",
           num_format=vfmt, **extra))
    lab = dict(bg_color=bg, font_color=C["mtd"], font_size=10, align="center", bold=True)
    l_b = dict(left=5, left_color=C["page"]) if gap_borders else {}
    r_b = dict(right=5, right_color=C["page"]) if gap_borders else {}
    text(ws, r0 + 4, c0, r0 + 4, mid - 1, "MTD", F(**lab, **l_b))
    text(ws, r0 + 4, mid, r0 + 4, c1, "MoM", F(**{**lab, "font_color": C["pink"]}, **r_b))
    text(ws, r0 + 5, c0, r0 + 5, mid - 1, f"={mtd}",
         F(bg_color=bg, font_color=C["text"], font_size=11, align="center", num_format=mfmt, **l_b))
    text(ws, r0 + 5, mid, r0 + 5, c1, f"={mom}",
         F(bg_color=bg, font_color=C["text"], font_size=11, align="center",
           num_format='0.0%;-0.0%;0.0%', **r_b))


# ----------------------------------------------------------------------------
# 1) Build slicer panels first (they define SEL cells), then calc, then visuals
# ----------------------------------------------------------------------------
for ws, page in ((ws_sum, "Summary"), (ws_ovr, "Overview")):
    setup_canvas(ws, page)
    text(ws, 8, 0, 8, 10, "Menu", F(bg_color=C["panel"], font_color=C["text"], bold=True,
                                   font_size=12, indent=1))
    nav_buttons(ws, page, [(9, 10), (12, 13), (15, 16)], 1, 9)

# Summary slicers: Purpose, Grade, State
r = 19
for s in PAGE_FILTERS["Summary"]:
    dropdown(ws_sum, "Summary", s, r, 1, 9, PANEL_LABEL, PANEL_DD)
    r += 3
text(ws_sum, r + 1, 1, r + 4, 9,
     "Tip: pick a value in any dropdown (or 'All') to slice every KPI, gauge and the grid.",
     PANEL_HINT)

# Overview slicers: Select Measure (field parameter) + Good v Bad, Grade, State
MEASURE_CELL = dropdown(ws_ovr, "Overview", "Select Measure", 18, 1, 9, PANEL_LABEL, PANEL_DD)
r = 21
for s in ["Good v Bad Loan", "Grade", "State"]:
    dropdown(ws_ovr, "Overview", s, r, 1, 9, PANEL_LABEL, PANEL_DD)
    r += 3
text(ws_ovr, 30, 0, 30, 10, "Cross-filter (chart click)",
     F(bg_color=C["panel"], font_color=C["mtd"], bold=True, font_size=10, indent=1))
r = 31
for s in ["Term", "Employee Length", "Purpose", "Home Ownership"]:
    dropdown(ws_ovr, "Overview", s, r, 1, 9,
             F(font_color=C["sub"], font_size=9, bg_color=C["panel"]), PANEL_DD)
    r += 3

# Details page uses its own column layout (wide table columns)
DET_W = [80, 160, 140, 90, 100, 150, 140, 100, 120, 130]
ws_det.hide_gridlines(2)
ws_det.set_zoom(90)
ws_det.set_column_pixels(0, 0, 220, F(bg_color=C["panel"]))
det_col_fmt = [
    dict(num_format="0"), dict(), dict(), dict(align="center"), dict(align="center"),
    dict(num_format="dd mmmm yyyy"), dict(num_format=NF["usd1"], font_color="#6FB6FF"),
    dict(num_format="0.00%"), dict(num_format="$#,##0.0"),
    dict(num_format=NF["usd1"], font_color=C["gold"]),
]
for i, w in enumerate(DET_W):
    ws_det.set_column_pixels(i + 1, i + 1, w, F(**{"bg_color": C["page"], "font_color": C["text"],
                                                   **det_col_fmt[i]}))
ws_det.set_column_pixels(11, 30, 20, F(bg_color=C["page"]))
for r in range(14):
    ws_det.set_row_pixels(r, 20)
header(ws_det, 0, 0, 2, 10, "Details")
text(ws_det, 3, 0, 3, 0, "Menu", F(bg_color=C["panel"], font_color=C["text"], bold=True,
                                   font_size=12, indent=1))
nav_buttons(ws_det, "Details", [(4, 5), (7, 8), (10, 11)], 0, 0)
DET_LBL = F(bg_color=C["page"], font_color=C["text"], bold=True, font_size=10, valign="bottom")
DET_DD = F(bg_color=C["dd"], font_color=C["text"], font_size=9, border=1,
           border_color="#9FB6E0", indent=1)
for s, (c0, c1) in zip(PAGE_FILTERS["Details"], [(1, 2), (3, 4), (5, 6)]):
    ws_det.merge_range(11, c0, 11, c1, s, DET_LBL)
    ws_det.merge_range(12, c0, 12, c1, "All", DET_DD)
    ws_det.data_validation(12, c0, 12, c1, {"validate": "list", "source": LIST_REF[s]})
    SEL["Details"][s] = absref("Details", 12, c0)

# ----------------------------------------------------------------------------
# 2) Calc: criteria, KPIs, breakdowns
# ----------------------------------------------------------------------------
calc.title("Slicer criteria (\"All\" -> wildcard)")
for page, slicers in PAGE_FILTERS.items():
    for s in slicers:
        CRIT[page][s] = calc.put(f"{page} | {s}",
                                 f'=IF({SEL[page][s]}="All","*","="&{SEL[page][s]})')

calc.title("Time intelligence (MTD = latest issue month in data, PMTD = month before)")
MTD = calc.put("MTD year-month key", "=MAX(d_ym)")
PMTD = calc.put("PMTD year-month key", f"=IF(MOD({MTD},100)=1,{MTD}-89,{MTD}-1)")
calc.put("MTD month", f'=DATE(INT({MTD}/100),MOD({MTD},100),1)', "mmm yyyy")
MEAS_IDX = calc.put("Selected measure # (field parameter)",
                    f"=MATCH({MEASURE_CELL},{LIST_REF['Select Measure'][1:]},0)")
MONEY = calc.put("Selected measure is money?", f"={MEAS_IDX}>1")
MEAS_NAME = calc.put("Selected measure name", f"={MEASURE_CELL}")


def kpi_block(page):
    calc.title(f"{page} KPIs")
    a = args(page)
    mtd = f"{a},d_ym,{MTD}"
    pmtd = f"{a},d_ym,{PMTD}"
    k = {}
    for key, label, agg, num in [
        ("apps", "Total Loan Applications", "COUNTIFS({})", NF["int"]),
        ("funded", "Total Funded Amount", "SUMIFS(d_amount,{})", NF["usd"]),
        ("recv", "Total Amount Received", "SUMIFS(d_payment,{})", NF["usd"]),
        ("rate", "Avg Interest Rate", "IFERROR(AVERAGEIFS(d_rate,{}),0)", NF["pct2"]),
        ("dti", "Avg DTI", "IFERROR(AVERAGEIFS(d_dti,{}),0)", NF["pct2"]),
    ]:
        k[key] = calc.put(label, "=" + agg.format(a), num)
        k[key + "_mtd"] = calc.put("  MTD " + label, "=" + agg.format(mtd), num)
        k[key + "_pmtd"] = calc.put("  PMTD " + label, "=" + agg.format(pmtd), num)
        k[key + "_mom"] = calc.put("  MoM % " + label,
                                   f'=IFERROR(({k[key + "_mtd"]}-{k[key + "_pmtd"]})/{k[key + "_pmtd"]},"")',
                                   NF["pct1"])
    return k


KPI = {p: kpi_block(p) for p in PAGE_FILTERS}

# Good vs Bad (Summary)
calc.title("Summary | Good loan vs Bad loan")
GB = {}
a = args("Summary")
for tag, crit in (("good", "Good Loan"), ("bad", "Bad Loan")):
    GB[tag + "_apps"] = calc.put(f"{crit} Applications", f'=COUNTIFS({a},d_gb,"{crit}")', NF["int"])
    GB[tag + "_pct"] = calc.put(f"{crit} %", f'=IFERROR({GB[tag + "_apps"]}/{KPI["Summary"]["apps"]},0)',
                                NF["pct1"])
    GB[tag + "_rest"] = calc.put(f"{crit} % remainder (gauge)", f'=1-{GB[tag + "_pct"]}', NF["pct1"])
    GB[tag + "_funded"] = calc.put(f"{crit} Funded Amount", f'=SUMIFS(d_amount,{a},d_gb,"{crit}")', NF["usd"])
    GB[tag + "_recv"] = calc.put(f"{crit} Total Received", f'=SUMIFS(d_payment,{a},d_gb,"{crit}")', NF["usd"])

# Loan status grid (Summary)
calc.title("Summary | Loan status grid")
calc.header(["Loan Status", "Total Loan Applications", "Total Funded Amount", "Total Amount Received",
             "MTD Funded Amount", "MTD Amount Received", "Avg Interest Rate", "Avg DTI"])
STATUS = ["Fully Paid", "Current", "Charged Off"]
GRID_START = calc.r
for s in STATUS + ["Grand Total"]:
    ex = [] if s == "Grand Total" else [("d_status", f'"{s}"')]
    base = args("Summary", extra=ex)
    mtd = base + f",d_ym,{MTD}"
    calc.ws.write(calc.r, 0, s)
    for j, (f, num) in enumerate([
        (f"=COUNTIFS({base})", NF["int"]), (f"=SUMIFS(d_amount,{base})", NF["usd"]),
        (f"=SUMIFS(d_payment,{base})", NF["usd"]), (f"=SUMIFS(d_amount,{mtd})", NF["usd"]),
        (f"=SUMIFS(d_payment,{mtd})", NF["usd"]),
        (f"=IFERROR(AVERAGEIFS(d_rate,{base}),0)", NF["pct2"]),
        (f"=IFERROR(AVERAGEIFS(d_dti,{base}),0)", NF["pct2"]),
    ], start=1):
        calc.ws.write_formula(calc.r, j, f, F(num_format=num))
    calc.r += 1


def breakdown(title, dim_label, cats, data_name, exclude, sort=False, label_kind="value"):
    """Overview breakdown table. Columns:
    A cat | B apps | C funded | D received | E selected | F label | G sort key
    and when sort=True: H rank | I sorted cat | J sorted value | K sorted label | L share
    """
    calc.title(title)
    cols = ["Category", "Applications", "Funded", "Received", "Selected measure", "Label", "Sort key"]
    if sort:
        cols += ["Rank", "Sorted category", "Sorted value", "Sorted label", "Share"]
    calc.header(cols)
    r0 = calc.r
    n = len(cats)
    rng = lambda c: f"Calc!${xl_col_to_name(c)}${r0 + 1}:${xl_col_to_name(c)}${r0 + n}"
    tot = f"SUM({rng(4)})"
    for i, cat in enumerate(cats):
        r = r0 + i
        crit_val = f'"={cat}"' if isinstance(cat, str) else str(cat)
        a = args("Overview", exclude=exclude, extra=[(data_name, crit_val)])
        calc.ws.write(r, 0, cat)
        calc.ws.write_formula(r, 1, f"=COUNTIFS({a})", F(num_format=NF["int"]))
        calc.ws.write_formula(r, 2, f"=SUMIFS(d_amount,{a})", F(num_format=NF["usd"]))
        calc.ws.write_formula(r, 3, f"=SUMIFS(d_payment,{a})", F(num_format=NF["usd"]))
        E = rc(r, 4)
        calc.ws.write_formula(r, 4, f"=CHOOSE({MEAS_IDX},{rc(r, 1)},{rc(r, 2)},{rc(r, 3)})",
                              F(num_format="#,##0"))
        lab = money_label(E, MONEY)
        if label_kind == "share":
            lab = f'{lab}&" ("&FIXED(IFERROR({E}/{tot},0)*100,1)&"%)"'
        calc.ws.write_formula(r, 5, "=" + lab)
        calc.ws.write_formula(r, 6, f"={E}+{(n - i)}/1000000")
        if sort:
            G = rng(6)
            calc.ws.write(r, 7, i + 1)
            calc.ws.write_formula(r, 9, f"=INDEX({rng(4)},MATCH(LARGE({G},{rc(r, 7)}),{G},0))",
                                  F(num_format="#,##0"))
            calc.ws.write_formula(r, 8, f"=INDEX({rng(0)},MATCH(LARGE({G},{rc(r, 7)}),{G},0))")
            calc.ws.write_formula(r, 11, f"=IFERROR({rc(r, 9)}/{tot},0)", F(num_format=NF["pct1"]))
            if label_kind == "tree":
                calc.ws.write_formula(
                    r, 10, f'=IF({rc(r, 11)}<0.035,"",{rc(r, 8)}&CHAR(10)&'
                           f'{money_label(rc(r, 9), MONEY)})')
            else:
                calc.ws.write_formula(r, 10, f"=INDEX({rng(5)},MATCH(LARGE({G},{rc(r, 7)}),{G},0))")
    calc.r += n
    ttl = calc.put("Chart title", f'={MEAS_NAME}&" by {dim_label}"')
    return dict(r0=r0, n=n, rng=rng, title=ttl)


B_MONTH = breakdown("Overview | Monthly trend by issue date", "Month", list(range(1, 13)),
                    "d_mon", None)
for i in range(12):
    calc.ws.write(B_MONTH["r0"] + i, 7, MONTHS[i])
B_STATE = breakdown("Overview | Regional analysis by state", "State", lists["State"],
                    "d_state", "State", sort=True)
B_TERM = breakdown("Overview | Loan term", "Term", lists["Term"], "d_term", "Term",
                   label_kind="share")
B_EMP = breakdown("Overview | Employee length", "Employee Length", lists["Employee Length"],
                  "d_emp", "Employee Length", sort=True)
B_PURP = breakdown("Overview | Loan purpose", "Purpose", lists["Purpose"],
                   "d_purpose", "Purpose", sort=True)
B_HOME = breakdown("Overview | Home ownership", "Home Ownership", lists["Home Ownership"],
                   "d_home", "Home Ownership", sort=True, label_kind="tree")

calc.title("Overview | Map helpers")
top5 = " & \"   ·   \" & ".join(
    f'INDEX({B_STATE["rng"](8)},{k})&" "&INDEX({B_STATE["rng"](10)},{k})' for k in range(1, 6))
TOP5 = calc.put("Top 5 states", f'="Top 5:  "&{top5}')

calc.title("Details")
DET_COUNT = calc.put("Rows shown", f'="Showing "&FIXED({KPI["Details"]["apps"]},0)&" loans"')

# ----------------------------------------------------------------------------
# Charts
# ----------------------------------------------------------------------------
WHITE = {"color": C["text"], "name": FONT, "size": 9}


def base_chart(kind, title_ref, subtype=None, bg=C["card"]):
    ch = wb.add_chart({"type": kind, **({"subtype": subtype} if subtype else {})})
    ch.set_chartarea({"border": {"none": True}, "fill": {"color": bg}})
    ch.set_plotarea({"border": {"none": True}, "fill": {"none": True}})
    if title_ref:
        ch.set_title({"name": "=" + title_ref, "overlay": False,
                      "name_font": {"color": C["sub"], "name": FONT, "size": 12, "bold": False}})
    else:
        ch.set_title({"none": True})
    ch.set_legend({"none": True})
    return ch


def custom_labels(b, col, n=None):
    n = n or b["n"]
    return [{"value": "=" + absref("Calc", b["r0"] + i, col)} for i in range(n)]


def gauge(ws, pct_ref, rest_ref, color, r, c, w, h):
    ch = wb.add_chart({"type": "doughnut"})
    ch.add_series({
        "values": f"={pct_ref}:{rest_ref.split('!')[1]}",
        "points": [{"fill": {"color": color}, "border": {"none": True}},
                   {"fill": {"color": C["grey"]}, "border": {"none": True}}],
    })
    ch.set_hole_size(80)
    ch.set_rotation(0)
    ch.set_chartarea({"border": {"none": True}, "fill": {"none": True}})
    ch.set_plotarea({"fill": {"none": True}})
    ch.set_title({"none": True})
    ch.set_legend({"none": True})
    ws.insert_chart(r, c, ch, {"x_scale": w / 480, "y_scale": h / 288, "object_position": 3})


# ----------------------------------------------------------------------------
# 3) SUMMARY dashboard
# ----------------------------------------------------------------------------
CARD_COLS = [(12, 23), (25, 36), (38, 49), (51, 62), (64, 75)]
KPI_DEFS = [
    ("Total Loan Applications", "apps", NF["k"], NF["k"]),
    ("Total Funded Amount", "funded", NF["m"], NF["m"]),
    ("Total Amount Received", "recv", NF["m"], NF["m"]),
    ("Avg Interest Rate", "rate", NF["pct2"], NF["pct1"]),
    ("Avg DTI", "dti", NF["pct2"], NF["pct1"]),
]


def kpi_row(ws, page, cols, gap=False):
    k = KPI[page]
    for i, ((c0, c1), (title, key, vf, mf)) in enumerate(zip(cols, KPI_DEFS)):
        kpi_card(ws, 4, c0, c1, title, (k[key], vf, k[key + "_mtd"], mf, k[key + "_mom"]),
                 hi=i == 0, gap_borders=gap)


kpi_row(ws_sum, "Summary", CARD_COLS)

for tag, c0, c1, color, title in (("good", 12, 43, C["gold"], "GOOD LOAN ISSUED"),
                                  ("bad", 45, 75, C["magenta"], "BAD LOAN ISSUED")):
    bg = C["card"]
    fill(ws_sum, 12, c0, 26, c1, F(bg_color=bg))
    text(ws_sum, 13, c0 + 1, 13, c0 + 13, title,
         F(bg_color=bg, font_color=C["blue"], bold=True, font_size=12, align="center"))
    gauge(ws_sum, GB[tag + "_pct"], GB[tag + "_rest"], color, 14, c0 + 1, 260, 240)
    val_color = C["gold"] if tag == "good" else "#8FC6FF"
    text(ws_sum, 19, c0 + 3, 20, c0 + 11, f"={GB[tag + '_pct']}",
         F(bg_color=bg, font_color=val_color, bold=True, font_size=24, align="center",
           num_format="0.0%"))
    bar_color = C["blue"] if tag == "good" else "#E8743B"
    for i, (lbl, key, nf) in enumerate([
        ("Applications", "_apps", NF["k"]), ("Funded Amount", "_funded", NF["m"]),
        ("Total Received", "_recv", NF["m"])]):
        rr = 14 + i * 4
        prefix = "Good Loan " if tag == "good" else "Bad Loan "
        text(ws_sum, rr, c0 + 16, rr, c1 - 1, prefix + lbl,
             F(bg_color=bg, font_color=C["text"], font_size=11, left=5, left_color=bar_color,
               indent=1))
        text(ws_sum, rr + 1, c0 + 16, rr + 2, c1 - 1, f"={GB[tag + key]}",
             F(bg_color=bg, font_color=val_color, bold=True, font_size=18, left=5,
               left_color=bar_color, indent=1, num_format=nf))

# Loan status grid
fill(ws_sum, 28, 12, 42, 75, F(bg_color=C["card"]))
text(ws_sum, 29, 13, 29, 30, "LOAN STATUS",
     F(bg_color=C["card"], font_color=C["blue"], bold=True, font_size=12))
GRID_COLS_SPAN = [(13, 19)] + [(20 + 8 * j, 27 + 8 * j) for j in range(7)]
heads = ["Loan Status", "Total Loan Applications", "Total Funded Amount", "Total Amount Received",
         "MTD Funded Amount", "MTD Amount Received", "Avg Interest Rate", "Avg DTI"]
for (a_, b_), h in zip(GRID_COLS_SPAN, heads):
    text(ws_sum, 30, a_, 31, b_, h, F(bg_color="#2B2B2B", font_color=C["text"], bold=True,
                                       font_size=10, align="center", text_wrap=True,
                                       bottom=2, bottom_color=C["blue"], right=1,
                                       right_color="#4A4A4A"))
grid_nf = [None, NF["int"], NF["usd"], NF["usd"], NF["usd"], NF["usd"], NF["pct2"], NF["pct2"]]
for i, s in enumerate(STATUS + ["Grand Total"]):
    rr = 32 + i * 2
    total = s == "Grand Total"
    bg = "#2B2B2B" if total else C["card"]
    for j, (a_, b_) in enumerate(GRID_COLS_SPAN):
        fmt = F(bg_color=bg, font_color=C["text"], font_size=11, bold=total or j == 0,
                align="left" if j == 0 else "right", indent=1, num_format=grid_nf[j] or "General",
                bottom=1, bottom_color="#29497F", right=1, right_color="#29497F")
        val = s if j == 0 else "=" + absref("Calc", GRID_START + i, j)
        text(ws_sum, rr, a_, rr + 1, b_, val, fmt)

# ----------------------------------------------------------------------------
# 4) OVERVIEW dashboard
# ----------------------------------------------------------------------------
kpi_row(ws_ovr, "Overview", CARD_COLS)


def cat_ref(b, col):
    return "=" + b["rng"](col)


# Monthly trend: area + line with labels
area = base_chart("area", B_MONTH["title"])
area.add_series({
    "categories": cat_ref(B_MONTH, 7), "values": cat_ref(B_MONTH, 4),
    "fill": {"color": C["area"], "transparency": 25}, "border": {"none": True},
})
line = wb.add_chart({"type": "line"})
line.add_series({
    "categories": cat_ref(B_MONTH, 7), "values": cat_ref(B_MONTH, 4),
    "line": {"color": C["line"], "width": 2.75},
    "marker": {"type": "circle", "size": 7, "fill": {"color": C["line"]},
               "border": {"color": C["line"]}},
    "data_labels": {"value": True, "position": "above", "font": {**WHITE, "bold": True},
                    "custom": custom_labels(B_MONTH, 5)},
})
area.combine(line)
area.set_x_axis({"num_font": WHITE, "line": {"none": True}, "major_tick_mark": "none",
                 "position_axis": "on_tick"})
area.set_y_axis({"visible": False, "major_gridlines": {"visible": False}})
ws_ovr.insert_chart(12, 12, area, {"x_scale": 480 / 480, "y_scale": 320 / 288, "object_position": 3})

# Regional analysis: tile map
fill(ws_ovr, 12, 37, 27, 58, F(bg_color=C["card"]))
text(ws_ovr, 12, 37, 13, 58, f"={B_STATE['title']}",
     F(bg_color=C["card"], font_color=C["sub"], font_size=12, indent=1))
MAP_R0, MAP_C0 = 15, 37
for st, (tc, tr) in TILES.items():
    r_, c_ = MAP_R0 + tr, MAP_C0 + tc * 2
    fmt = F(bg_color=C["map_na"], font_color="#1B1B1B", font_size=8, bold=True, align="center",
            border=1, border_color=C["card"], num_format=f'"{st}";"{st}";"{st}";"{st}"')
    formula = (f'=IFERROR(INDEX({B_STATE["rng"](4)},MATCH("{st}",{B_STATE["rng"](0)},0)),"")'
               if st in lists["State"] else '=""')
    ws_ovr.merge_range(r_, c_, r_, c_ + 1, "", fmt)
    ws_ovr.write_formula(r_, c_, formula, fmt)
ws_ovr.conditional_format(MAP_R0, MAP_C0, MAP_R0 + 7, MAP_C0 + 21, {
    "type": "2_color_scale", "min_color": C["map_lo"], "max_color": C["map_hi"]})
text(ws_ovr, 24, 38, 25, 57, f"={TOP5}",
     F(bg_color=C["card"], font_color=C["text"], font_size=9, text_wrap=True, align="center"))
text(ws_ovr, 26, 38, 26, 57, "Lighter = lower, darker = higher.  Grey = no loans.  "
                             "Use the State slicer to drill in.",
     F(bg_color=C["card"], font_color=C["sub"], font_size=8, italic=True, align="center"))

# Loan term donut
donut = base_chart("doughnut", B_TERM["title"])
donut.add_series({
    "categories": cat_ref(B_TERM, 0), "values": cat_ref(B_TERM, 4),
    "points": [{"fill": {"color": C["t36"]}, "border": {"color": C["card"]}},
               {"fill": {"color": C["t60"]}, "border": {"color": C["card"]}}],
    "data_labels": {"value": True, "font": {**WHITE, "bold": True, "color": "#0B0B0B"},
                    "custom": custom_labels(B_TERM, 5)},
})
donut.set_hole_size(55)
donut.set_legend({"position": "top", "font": WHITE})
ws_ovr.insert_chart(12, 60, donut, {"x_scale": 320 / 480, "y_scale": 320 / 288, "object_position": 3})


def hbar(b, color, r, c, w, h):
    ch = base_chart("bar", b["title"])
    ch.add_series({
        "categories": cat_ref(b, 8), "values": cat_ref(b, 9),
        "fill": {"color": color}, "border": {"none": True}, "gap": 35,
        "data_labels": {"value": True, "position": "outside_end", "font": WHITE,
                        "custom": custom_labels(b, 10)},
    })
    ch.set_y_axis({"reverse": True, "num_font": WHITE, "line": {"none": True},
                   "major_tick_mark": "none"})
    ch.set_x_axis({"visible": False, "major_gridlines": {"visible": False}})
    ws_ovr.insert_chart(r, c, ch, {"x_scale": w / 480, "y_scale": h / 288, "object_position": 3})


hbar(B_EMP, C["lav"], 29, 12, 480, 300)
hbar(B_PURP, C["purp"], 29, 37, 440, 300)

# Home ownership "treemap": single 100% stacked column, one block per category (largest on top)
tree = base_chart("column", B_HOME["title"], subtype="percent_stacked")
for k in reversed(range(B_HOME["n"])):
    rr = B_HOME["r0"] + k
    tree.add_series({
        "name": "=" + absref("Calc", rr, 8),
        "values": "=" + absref("Calc", rr, 9),
        "fill": {"color": HOME_COLORS[k % len(HOME_COLORS)]},
        "border": {"color": C["card"], "width": 1.5}, "gap": 4, "overlap": 100,
        "data_labels": {"value": True, "position": "inside_base",
                        "font": {**WHITE, "size": 11, "bold": True},
                        "custom": [{"value": "=" + absref("Calc", rr, 10)}]},
    })
tree.set_x_axis({"visible": False})
tree.set_y_axis({"visible": False, "major_gridlines": {"visible": False}})
ws_ovr.insert_chart(29, 60, tree, {"x_scale": 320 / 480, "y_scale": 300 / 288, "object_position": 3})

# ----------------------------------------------------------------------------
# 5) DETAILS dashboard
# ----------------------------------------------------------------------------
fill(ws_det, 3, 1, 3, 10, F(bg_color=C["page"]))
kpi_row(ws_det, "Details", [(1 + 2 * i, 2 + 2 * i) for i in range(5)], gap=True)
fill(ws_det, 11, 7, 11, 10, F(bg_color=C["page"]))
text(ws_det, 12, 8, 12, 10, f"={DET_COUNT}",
     F(bg_color=C["page"], font_color=C["sub"], italic=True, align="right"))
ws_det.write_blank(12, 7, None, F(bg_color=C["page"]))
for j, h in enumerate(["id", "purpose", "home_ownership", "grade", "sub_grade", "issue_date",
                       "Funded Amount", "Int Rate", "Installment", "Amount Collection"]):
    ws_det.write(13, j + 1, h, F(bg_color="#1C1C1C", font_color=C["text"], bold=True,
                                 align="center", bottom=2, bottom_color=C["blue"]))
conds = "*".join(f'(({SEL["Details"][s]}="All")+({DIM[s]}={SEL["Details"][s]}))'
                 for s in PAGE_FILTERS["Details"])
ws_det.write_dynamic_array_formula(
    14, 1, 14, 1,
    f'=FILTER(Data!$A$2:$J${LAST},{conds},"No loans match the selected filters")')
ws_det.conditional_format(14, 1, 14 + N, 10, {
    "type": "formula", "criteria": "=AND($B15<>\"\",MOD(ROW(),2)=0)",
    "format": wb.add_format({"bg_color": C["row1"]})})
ws_det.conditional_format(14, 1, 14 + N, 10, {
    "type": "formula", "criteria": "=AND($B15<>\"\",MOD(ROW(),2)=1)",
    "format": wb.add_format({"bg_color": C["row2"]})})
ws_det.freeze_panes(14, 1)

# ----------------------------------------------------------------------------
# 6) Guide sheet
# ----------------------------------------------------------------------------
g = ws_guide
g.hide_gridlines(2)
g.set_column(0, 0, 4)
g.set_column(1, 1, 34)
g.set_column(2, 2, 110)
h = F(bold=True, font_size=16, font_color="#0B2D5E")
hh = F(bold=True, font_size=11, bg_color="#0B2D5E", font_color="#FFFFFF")
t = F(text_wrap=True, valign="top")
tb = F(text_wrap=True, valign="top", bold=True)
g.write(1, 1, "Bank Loan Report - Excel edition (how it maps to the Power BI project)", h)
rows_ = [
    ("PAGE / FEATURE", "HOW IT IS BUILT IN EXCEL"),
    ("Summary page", "5 KPI cards (Total / MTD / MoM), Good-vs-Bad loan gauges + KPIs, "
                     "Loan Status grid. Slicers: Purpose, Grade, State."),
    ("Overview page", "Same KPI cards + 6 charts: Monthly trend (area+line), Regional map "
                      "(US tile map, heat-coloured), Loan Term donut, Employee Length bars, "
                      "Purpose bars, Home Ownership treemap-style blocks. Slicers: Select Measure, "
                      "Good v Bad Loan, Grade, State."),
    ("Details page", "KPI cards + full loan grid (id, purpose, home ownership, grade, sub grade, "
                     "issue date, funded amount, int rate, installment, amount collected) built "
                     "with one FILTER() spill. Slicers: Good v Bad Loan, Grade, State. Header and "
                     "KPIs are frozen so the grid scrolls like the Power BI table."),
    ("Navigation buttons", "Summary / Overview / Details buttons on every page are hyperlinks; "
                           "the active page's button is highlighted."),
    ("Slicers", "Dropdown cells (Data Validation). 'All' = no filter. Every KPI/chart is a "
                "COUNTIFS / SUMIFS / AVERAGEIFS on the Data sheet, so changes are instant."),
    ("Field parameter (Select Measure)", "Overview dropdown switches every chart, its labels and "
                                          "its title between Total Loan Applications, Total Funded "
                                          "Amount and Total Amount Received (CHOOSE on the measure #)."),
    ("Filter interactions (click a bar)", "Excel charts cannot cross-filter on click, so the "
                                          "Overview 'Cross-filter' dropdowns (Term, Employee "
                                          "Length, Purpose, Home Ownership) do the same job: they "
                                          "filter the KPIs and every other chart, while the chart "
                                          "of that dimension keeps showing all its categories."),
    ("Good vs Bad grouping", "Data column good_vs_bad: Fully Paid + Current = Good Loan, "
                             "Charged Off = Bad Loan (same as the Power BI group)."),
    ("MTD / PMTD / MoM", "MTD = latest issue month in the data (Dec 2021), PMTD = previous month; "
                         "MoM % = (MTD - PMTD) / PMTD. Matches TOTALMTD / DATEADD in the video."),
    ("Sorting", "Employee Length, Purpose, State and Home Ownership are ranked by the selected "
                "measure with LARGE/INDEX/MATCH, so bars stay sorted after any filter."),
    ("Calc sheet", "All intermediate formulas, readable top to bottom (criteria, KPIs, grids, "
                   "chart series and labels)."),
    ("Data sheet", "38,576 loans as the Excel table tblLoan, with helper columns good_vs_bad, "
                   "issue_month, issue_month_name, year_month. Named ranges d_* point at it."),
    ("Requirements", "Excel 2021 or Microsoft 365 (the Details grid uses FILTER). Everything else "
                     "works in Excel 2010+. No macros, no add-ins, no external connections."),
    ("Refreshing data", "Replace data/financial_loan.csv and run: python build_dashboard.py"),
]
for i, (a_, b_) in enumerate(rows_):
    r = 3 + i
    if i == 0:
        g.write(r, 1, a_, hh)
        g.write(r, 2, b_, hh)
    else:
        g.write(r, 1, a_, tb)
        g.write(r, 2, b_, t)
        g.set_row(r, 32)

ws_calc.set_tab_color("#7F7F7F")
ws_list.set_tab_color("#7F7F7F")
ws_data.set_tab_color("#7F7F7F")
for ws, color in ((ws_sum, "#3DA5FF"), (ws_ovr, "#3DA5FF"), (ws_det, "#3DA5FF")):
    ws.set_tab_color(color)
for ws in (ws_sum, ws_ovr):
    ws.set_landscape()
    ws.set_paper(9)
    ws.set_margins(0.2, 0.2, 0.2, 0.2)
    ws.print_area(0, 0, GRID_ROWS - 1, GRID_COLS - 1)
    ws.fit_to_pages(1, 1)
ws_det.set_landscape()
ws_det.set_paper(9)
ws_det.set_margins(0.2, 0.2, 0.2, 0.2)
ws_det.print_area(0, 0, 14 + N, 10)
ws_det.repeat_rows(13)
ws_det.fit_to_pages(1, 0)

ws_list.hide()
ws_sum.activate()
wb.close()
print("wrote", OUT)
